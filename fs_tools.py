from pathlib import Path
from datetime import datetime, timezone
from mimetypes import guess_type
from docx import Document
from pypdf import PdfReader

READABLE_EXTENSIONS = {".txt", ".pdf", ".docx"}


def read_txt(path):
    return Path(path).read_text(encoding="utf-8")


def read_pdf(path):
    reader = PdfReader(path)
    text_parts = []

    for page in reader.pages:
        text = page.extract_text()
        if text:
            text_parts.append(text)

    return " ".join(text_parts)


def read_docx(path):
    document = Document(path)
    paragraphs = []

    for paragraph in document.paragraphs:
        paragraphs.append(paragraph.text)

    return " ".join(paragraphs)


def write_txt(path, content):
    Path(path).write_text(str(content), encoding="utf-8")


def write_docx(path, content):
    document = Document()

    lines = str(content).splitlines()
    if not lines:
        document.add_paragraph("")
    else:
        for line in lines:
            document.add_paragraph(line)

    document.save(path)


def is_word_character(character):
    return character.isalnum() or character == "_"


def expand_start_to_word_boundary(text, index):
    while (
        index > 0
        and is_word_character(text[index])
        and is_word_character(text[index - 1])
    ):
        index -= 1

    return index


def expand_end_to_word_boundary(text, index):
    while (
        index > 0
        and index < len(text)
        and is_word_character(text[index - 1])
        and is_word_character(text[index])
    ):
        index += 1

    return index


def get_file_metadata(path):
    path = Path(path)
    stats = path.stat()
    mime_type, encoding = guess_type(path)

    return {
        "name": path.name,
        "stem": path.stem,
        "extension": path.suffix.lower(),
        "path": str(path),
        "absolute_path": str(path.resolve()),
        "size_bytes": stats.st_size,
        "created_at": datetime.fromtimestamp(stats.st_ctime, tz=timezone.utc).isoformat(),
        "modified_at": datetime.fromtimestamp(stats.st_mtime, tz=timezone.utc).isoformat(),
        "mime_type": mime_type,
        "encoding": encoding
    }

def read_file(path):
    path = Path(path)
    extension = path.suffix.lower()

    if extension == ".txt":
        text = read_txt(path)
        reader = "read_txt"

    elif extension == ".pdf":
        text = read_pdf(path)
        reader = "read_pdf"

    elif extension == ".docx":
        text = read_docx(path)
        reader = "read_docx"

    else:
        raise ValueError(f"Unsupported file type: {extension}")

    return {
        "text": text,
        "metadata": get_file_metadata(path),
        "file_type": extension.lstrip("."),
        "reader": reader,
        "text_length": len(text),
        "word_count": len(text.split()),
        "line_count": len(text.splitlines()),
        "is_empty": len(text.strip()) == 0,
    }


def write_file(path, content, overwrite=True, create_parent_dirs=True):
    path = Path(path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        return {
            "message": "The file is a PDF and cannot be edited",
            "file_type": "pdf",
            "written": False,
        }

    if path.exists() and path.is_dir():
        raise IsADirectoryError(f"Cannot write to a directory: {path}")

    if path.exists() and not overwrite:
        raise FileExistsError(f"File already exists: {path}")

    if create_parent_dirs:
        path.parent.mkdir(parents=True, exist_ok=True)

    if extension == ".txt":
        write_txt(path, content)
        writer = "write_txt"

    elif extension == ".docx":
        write_docx(path, content)
        writer = "write_docx"

    else:
        raise ValueError(f"Unsupported file type: {extension}")

    return {
        "message": "File written successfully",
        "metadata": get_file_metadata(path),
        "file_type": extension.lstrip("."),
        "writer": writer,
        "characters_written": len(str(content)),
        "line_count": len(str(content).splitlines()),
    }


def search_in_file(filepath, keyword):
    if not keyword:
        raise ValueError("Keyword cannot be empty")

    file_data = read_file(filepath)
    text = file_data["text"]
    keyword_lower = keyword.lower()
    text_lower = text.lower()
    context_size = 50
    matches = []
    start_at = 0

    while True:
        match_start = text_lower.find(keyword_lower, start_at)
        if match_start == -1:
            break

        match_end = match_start + len(keyword)
        context_start = max(0, match_start - context_size)
        context_end = min(len(text), match_end + context_size)
        context_start = expand_start_to_word_boundary(text, context_start)
        context_end = expand_end_to_word_boundary(text, context_end)
        pre_context = text[context_start:match_start]
        matched_text = text[match_start:match_end]
        post_context = text[match_end:context_end]

        matches.append({
            "pre_context": pre_context,
            "match": matched_text,
            "post_context": post_context,
            "full_match": pre_context + matched_text + post_context,
            "start_index": match_start,
            "end_index": match_end,
        })

        start_at = match_end

    return {
        "keyword": keyword,
        "filepath": str(filepath),
        "metadata": file_data["metadata"],
        "match_count": len(matches),
        "matches": matches,
    }


def search_in_directory(directory, keyword):
    directory = Path(directory)

    if not directory.exists():
        raise FileNotFoundError(f"Directory not found: {directory}")

    if not directory.is_dir():
        raise NotADirectoryError(f"Not a directory: {directory}")

    results = []
    searched_files = 0
    skipped_files = []
    errors = []

    for path in sorted(directory.iterdir()):
        if not path.is_file():
            continue

        if path.suffix.lower() not in READABLE_EXTENSIONS:
            skipped_files.append(str(path))
            continue

        searched_files += 1

        try:
            result = search_in_file(path, keyword)
        except Exception as error:
            errors.append({
                "filepath": str(path),
                "error": str(error),
                "error_type": type(error).__name__,
            })
            continue

        if result["match_count"] > 0:
            results.append(result)

    total_matches = sum(result["match_count"] for result in results)

    return {
        "keyword": keyword,
        "directory": str(directory),
        "searched_files": searched_files,
        "matched_files": len(results),
        "total_matches": total_matches,
        "results": results,
        "skipped_files": skipped_files,
        "errors": errors,
    }


def list_files(directory, extension=None):
    directory = Path(directory)

    if not directory.exists():
        raise FileNotFoundError(f"Directory not found: {directory}")

    if not directory.is_dir():
        raise NotADirectoryError(f"Not a directory: {directory}")

    if extension:
        extension = extension.lower()
        if not extension.startswith("."):
            extension = f".{extension}"

    files = []

    for path in sorted(directory.iterdir()):
        if not path.is_file():
            continue

        if extension and path.suffix.lower() != extension:
            continue

        files.append(path.name)

    return files
